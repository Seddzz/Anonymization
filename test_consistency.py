#!/usr/bin/env python3
"""
Test script to verify DOCX anonymization consistency
"""

from docx import Document
from src.agent.executor import AnonymizerPipeline
from src.utils.document_processor import DocumentProcessor
import os

def create_test_docx():
    """Create a simple test DOCX file"""
    doc = Document()
    
    # Add test content with obvious entities
    doc.add_paragraph("Hello, my name is John Smith and I work at Microsoft Corporation in Seattle.")
    doc.add_paragraph("You can reach me at john.smith@microsoft.com or call me at +1-206-555-0123.")
    doc.add_paragraph("I am 34 years old and live at 123 Main Street, Seattle, WA 98101.")
    doc.add_paragraph("My colleague Sarah Johnson also works here, she's 27 years old.")
    doc.add_paragraph("My social security number is 123-45-6789.")
    
    test_file = "test_consistency.docx"
    doc.save(test_file)
    print(f"✅ Created test file: {test_file}")
    return test_file

def test_anonymization_consistency():
    """Test that anonymization is consistent between UI and document"""
    
    # Create test file
    test_file = create_test_docx()
    
    try:
        # Create pipeline and processor
        pipeline = AnonymizerPipeline(detector_type='spacy')
        processor = DocumentProcessor(pipeline)
        
        print("\n🔍 Testing DOCX anonymization...")
        
        # Process the file
        result = processor.process_file(test_file, 'docx')
        
        if result and result.get('success'):
            print("✅ Processing successful!")
            print(f"📄 Output file: {result.get('output_path')}")
            print(f"🔄 Replacements found: {len(result.get('replacement_mapping', {}))}")
            
            # Display the replacement mapping
            replacement_mapping = result.get('replacement_mapping', {})
            print("\n📋 Replacement Mapping:")
            for original, replacement in replacement_mapping.items():
                print(f"  '{original}' → '{replacement}'")
            
            # Now read the actual anonymized document to verify consistency
            if result.get('output_path') and os.path.exists(result['output_path']):
                anonymized_doc = Document(result['output_path'])
                anonymized_text = "\n".join([p.text for p in anonymized_doc.paragraphs if p.text.strip()])
                
                print(f"\n📖 Anonymized document content (first 300 chars):")
                print(f"'{anonymized_text[:300]}...'")
                
                # Check if the replacements in the mapping actually appear in the document
                consistency_check = True
                for original, replacement in replacement_mapping.items():
                    if original in anonymized_text:
                        print(f"❌ INCONSISTENCY: Original text '{original}' still found in document!")
                        consistency_check = False
                    elif replacement not in anonymized_text:
                        print(f"❌ INCONSISTENCY: Replacement '{replacement}' not found in document!")
                        consistency_check = False
                
                if consistency_check:
                    print("✅ CONSISTENCY CHECK PASSED: All replacements match between UI and document!")
                else:
                    print("❌ CONSISTENCY CHECK FAILED: Mismatches found!")
                    
            else:
                print("❌ Output file not found")
                
        else:
            print("❌ Processing failed:")
            print(f"Error: {result.get('error', 'Unknown error')}")
            
    except Exception as e:
        print(f"❌ Test failed with exception: {e}")
        import traceback
        traceback.print_exc()
    
    finally:
        # Clean up
        for file in [test_file, test_file.replace('.docx', '_anonymized.docx')]:
            if os.path.exists(file):
                os.remove(file)
                print(f"🧹 Cleaned up: {file}")

if __name__ == "__main__":
    test_anonymization_consistency()